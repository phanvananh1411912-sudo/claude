import json, glob, re, sys
from openpyxl import load_workbook

SP = '/tmp/claude-0/-home-user-claude/8d38f015-ed38-5363-ae81-d0cc7e396864/scratchpad'
TEMPLATE = '/root/.claude/uploads/8d38f015-ed38-5363-ae81-d0cc7e396864/e14db25e-MAU_Trac_nghiem_bai_hoc.xlsx'
OUT = '/home/user/claude/hsk-data/HSK4NEXUS-CauHoiTracNghiem-200bai.xlsx'

data = json.load(open(f'{SP}/raw_quiz.json', encoding='utf-8'))

# fix known source corruption (mojibake character), preserving meaning
for q in data:
    if q['bai'] == 'HSK2-Bài 27' and q['so'] == 3:
        if q['options']['B']['han'] == '心里有点�一想放弃':
            q['options']['B']['han'] = '心里有点儿想放弃'
            print('Đã sửa ký tự lỗi: HSK2-Bài 27 câu 3 đáp án B')

enr = {}
for f in glob.glob(f'{SP}/quiz_batches/*.out.json'):
    for item in json.load(open(f, encoding='utf-8')):
        enr[(item['bai'], item['so'])] = item

def bai_no(b):
    return int(re.search(r'\d+', b).group())

def group_rank(b):
    if b.startswith('HSK3-Hội thoại'): return 4
    if b.startswith('HSK3-Bài'): return 3
    if b.startswith('HSK2-Hội thoại'): return 2
    if b.startswith('HSK2-Bài'): return 1
    return 0

KIN_FIX = [
    (r'\bbà ba\b', 'bàba'), (r'\bBà ba\b', 'Bàba'),
    (r'\bmā ma\b', 'māma'), (r'\bMā ma\b', 'Māma'),
    (r'\bgē ge\b', 'gēge'), (r'\bGē ge\b', 'Gēge'),
    (r'\bdì di\b', 'dìdi'), (r'\bDì di\b', 'Dìdi'),
    (r'\bmèi mei\b', 'mèimei'), (r'\bMèi mei\b', 'Mèimei'),
    (r'\byé ye\b', 'yéye'), (r'\bYé ye\b', 'Yéye'),
    (r'\bnǎi nai\b', 'nǎinai'), (r'\bNǎi nai\b', 'Nǎinai'),
]
def fix_pinyin(s):
    for pat, rep in KIN_FIX:
        s = re.sub(pat, rep, s)
    return s

errors, warns = [], []
rows = []
for q in sorted(data, key=lambda q: (group_rank(q['bai']), bai_no(q['bai']), q['so'])):
    key = (q['bai'], q['so'])
    e = enr.get(key)
    if e is None:
        errors.append(f'MISSING enrichment {key}'); continue
    cau_hoi_py = fix_pinyin(e.get('cau_hoi_pinyin', '').strip())
    cau_hoi_vi = e.get('cau_hoi_vi', '').strip()
    if not cau_hoi_py or not cau_hoi_vi:
        errors.append(f'{key}: thiếu cau_hoi pinyin/vi')
    row = [q['bai'], q['so'], q['cau_hoi'], cau_hoi_py, cau_hoi_vi]
    is_hsk1_inline = q['bai'].startswith('Bài') and bai_no(q['bai']) <= 20
    for k in 'ABCD':
        o = q['options'][k]
        if is_hsk1_inline and o.get('pinyin'):
            py, vi = fix_pinyin(o['pinyin'].strip()), o['vi'].strip()
        else:
            eo = e.get('options', {}).get(k, {})
            py, vi = fix_pinyin(eo.get('pinyin', '').strip()), eo.get('vi', '').strip()
        if not py or not vi:
            errors.append(f'{key} opt{k}: thiếu pinyin/vi')
        row += [o['han'], py, vi]
    row.append(q['dap_an'])
    assert len(row) == 18
    rows.append(row)

print(f'Tổng số dòng: {len(rows)} | ERRORS: {len(errors)}')
for e in errors[:30]: print(' E:', e)
if errors:
    sys.exit(1)

wb = load_workbook(TEMPLATE)
ws = wb['Câu hỏi']
# xóa các dòng mẫu minh họa (giữ header dòng 1)
ws.delete_rows(2, ws.max_row - 1)
for row in rows:
    ws.append(row)
wb.save(OUT)
print('built', OUT, '| tổng dòng dữ liệu:', ws.max_row - 1)
