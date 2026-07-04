"""Builder tạo file Excel theo template HSK4 NEXUS (3 sheet: Câu, Từ vựng, Ngữ pháp).

Dùng: from hsk_excel_builder import build_workbook
build_workbook(path, cau_rows, tuvung_rows, nguphap_rows)
  cau_rows:     list[(bai, stt, hanzi, pinyin, nghia, nguphap)]
  tuvung_rows:  list[(bai, hanzi, pinyin, hanviet, pos, meaning)]
  nguphap_rows: list[(bai, cautruc, ynghia)]  # Số lần=0, Mục tiêu=10 tự thêm
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADER_FILL = PatternFill('solid', start_color='C00000')
HEADER_FONT = Font(bold=True, color='FFFFFF', name='Arial')
CENTER = Alignment(horizontal='center')
BODY_FONT = Font(name='Arial')

SHEETS = {
    'Câu': (['Bài', '#', 'Hán tự', 'Pinyin', 'Nghĩa', 'Ngữ pháp'],
            [14, 5, 45, 55, 55, 30]),
    'Từ vựng': (['Bài', 'Hanzi', 'Pinyin', 'HanViet', 'POS', 'Meaning'],
                [14, 14, 18, 18, 8, 40]),
    'Ngữ pháp': (['Bài', 'Cấu trúc', 'Ý nghĩa', 'Số lần', 'Mục tiêu'],
                 [14, 35, 60, 8, 10]),
}


def build_workbook(path, cau_rows, tuvung_rows, nguphap_rows):
    wb = Workbook()
    wb.remove(wb.active)
    data = {
        'Câu': cau_rows,
        'Từ vựng': tuvung_rows,
        'Ngữ pháp': [(b, ct, yn, 0, 10) for (b, ct, yn) in nguphap_rows],
    }
    for name, (headers, widths) in SHEETS.items():
        ws = wb.create_sheet(name)
        for col, (h, w) in enumerate(zip(headers, widths), 1):
            c = ws.cell(row=1, column=col, value=h)
            c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, CENTER
            ws.column_dimensions[c.column_letter].width = w
        for r, row in enumerate(data[name], 2):
            for col, v in enumerate(row, 1):
                ws.cell(row=r, column=col, value=v).font = BODY_FONT
        ws.freeze_panes = 'A2'
    wb.save(path)
    return path
