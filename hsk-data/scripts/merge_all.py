import sys, os
sys.path.insert(0, '/tmp/claude-0/-home-user-claude/8d38f015-ed38-5363-ae81-d0cc7e396864/scratchpad')
from openpyxl import load_workbook
from hsk_excel_builder import build_workbook

HD = '/home/user/claude/hsk-data'
FILES = [
    'HSK4NEXUS-TronGoi-120bai.xlsx',       # HSK1 40 + HSK2 doc 40 + HSK2 hoi thoai 40 = 120
    'HSK4NEXUS-HSK3-40bai.xlsx',           # HSK3 doc 40
    'HSK4NEXUS-HSK3-HoiThoai-6bai.xlsx',   # HSK3 hoi thoai 1-6
    'HSK4NEXUS-HSK3-HoiThoai-Bai7-13.xlsx',
    'HSK4NEXUS-HSK3-HoiThoai-Bai14-22.xlsx',
    'HSK4NEXUS-HSK3-HoiThoai-Bai23-40.xlsx',
]

cau, tv, np_ = [], [], []
seen_bai_cau = set()
for fn in FILES:
    wb = load_workbook(f'{HD}/{fn}')
    for row in wb['Câu'].iter_rows(min_row=2, values_only=True):
        if row[0] is None: continue
        cau.append(row)
        seen_bai_cau.add(row[0])
    for row in wb['Từ vựng'].iter_rows(min_row=2, values_only=True):
        if row[0] is None: continue
        tv.append(row)
    for row in wb['Ngữ pháp'].iter_rows(min_row=2, values_only=True):
        if row[0] is None: continue
        np_.append((row[0], row[1], row[2]))  # drop Số lần/Mục tiêu, builder re-adds them

print(f'Tổng số bài (sheet Câu): {len(seen_bai_cau)}')
print(f'Câu: {len(cau)} | Từ vựng: {len(tv)} | Ngữ pháp: {len(np_)}')

out = f'{HD}/HSK4NEXUS-TatCa-HSK1-2-3.xlsx'
build_workbook(out, cau, tv, np_)
print('built', out)
