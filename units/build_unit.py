# -*- coding: utf-8 -*-
"""Tạo file Unit (khoá Tiếng Anh Mất Căn Bản) từ 5 câu bài C10T2 - Gifted children and learning."""
from docx import Document
from docx.shared import Pt, Inches, Twips
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Cambria"
TABLE_PT = 10.0
MIN_PT = 6.5
DXA_PER_CHAR = 135
MIN_COL = 800
PAGE_W = 16838  # A4 landscape, twips
MARGIN = 576    # 0.4in
USABLE = PAGE_W - 2 * MARGIN

# Nhãn chức năng ngữ pháp -> in đậm ở hàng 2
BOLD_TOKENS = {
    "S", "TT", "BC", "DV", "Trạng ngữ", "Tính ngữ", "MĐC", "MĐP", "MĐQH",
    "CNTT", "Conj", "Conjp", "Aux", "Advpđ", "ĐX", "Vlet",
    "Vi", "Vt", "Vikr", "Vtpd", "Vt[bđ]",
}

KY_HIEU_1 = (
    "Chức năng ngữ pháp: S=Chủ ngữ ; TT=Tân trực tiếp ; BC=Bổ chủ ; DV=Đồng vị ngữ ; "
    "Trạng ngữ=Adv / Adv+N[...] / Prep+N[...] ; Tính ngữ=Prep+N[...] sau danh từ ; "
    "MĐC=Mệnh đề chính ; MĐP=Mệnh đề phụ (câu phụ) ; MĐQH=Mệnh đề quan hệ (relative clause) ; "
    "CNTT=Chủ ngữ thật (trong câu chủ ngữ giả “it”) ; Conj=Liên từ đẳng lập (and, but, or, while...) ; "
    "Conjp=Liên từ phụ thuộc (because, when, if, although, so that...) ; "
    "Aux=Trợ động từ (will, can, may, must, should, do, have...) ; "
    "Advpđ=Trạng từ phủ định (not, never, no longer, hardly...) ; "
    "ĐX=Thành phần đệm-xen (thán từ, hô ngữ, câu chêm: Oh, Come on, sir, Excuse me...) ; "
    "Vlet=Cấu trúc đề nghị “Let's” (Let us) ; "
    "Vt=Ngoại động từ ; Vtpd=Động từ phổ dụng ; Vi=Nội động từ ; Vikr=Động từ nối (linking verb) ; "
    "Vt[bđ]=Ngoại động từ ở thể bị động (be + V3/V-ed)"
)

KY_HIEU_2 = (
    "Loại từ: Pro=Đại từ nhân xưng / đại từ quan hệ ; Pro (CN giả)=Đại từ chủ ngữ giả/hình thức (“it”, “there” không mang nghĩa) ; "
    "PN=Danh từ riêng ; N[C,S]=Danh từ đếm được số ít ; N[C,P]=Danh từ đếm được số nhiều ; "
    "N[U]=Danh từ không đếm được ; N1N2=Danh từ kép ; Article=Mạo từ (a/an/the) ; "
    "Số từ=Số đếm (two, three...) ; Adjsh=Tính từ sở hữu ; Adjcd=Tính từ chỉ định ; "
    "Adjbd=Tính từ bất định (less, some, many...) ; Adj=Tính từ miêu tả ; "
    "Adv=Trạng từ (bổ nghĩa cho tính từ/động từ: very, really...) ; Prep=Giới từ ; "
    "Adv[N]=Danh từ chỉ thời gian (day, night, month, year, hour, morning...) dùng như trạng từ — "
    "khi đi với every/per, HOẶC khi cả cụm trạng ngữ không có giới từ đứng trước ; "
    "Adv[Adj]=Tính từ dùng như trạng ngữ ; V-ing=Danh động từ (gerund) ; "
    "to V=Cụm động từ nguyên mẫu có “to” ; "
    "V=Động từ nguyên mẫu không “to” (bare infinitive, sau help/make/let...) ; "
    "how to V=Cụm “how + to V” (cách làm gì)"
)

SENTENCES = [
    {
        "num": "Câu 1",
        "vn": "Với những học sinh giỏi có khả năng làm bài chính xác, có thể rút ngắn thời gian làm bài tập.",
        "en": "Less time can be spent on exercises with gifted pupils who produce accurate work.",
        "cols": [
            ("Less [time]", "Adjbd+N[U]=S"),
            ("[can]", "Aux"),
            ("[be spent]", "Vt[bđ]"),
            ("[on exercises]", "Trạng ngữ=Prep+N[C,P]"),
            ("[with gifted pupils]", "Trạng ngữ=Prep+Adj+N[C,P]"),
            ("[who produce accurate work].", "MĐQH=Pro+Vt+TT=Adj+N[U]"),
        ],
    },
    {
        "num": "Câu 2",
        "vn": "Tự lực là một công cụ quý giá giúp học sinh có năng khiếu đạt được mục tiêu của mình.",
        "en": "Self-reliance is a valuable tool that helps gifted students reach their goals.",
        "cols": [
            ("[Self-reliance]", "N[U]=S"),
            ("[is]", "Vikr"),
            ("a valuable [tool]", "Article+Adj+N[C,S]=BC"),
            ("[that helps gifted students reach their goals].",
             "MĐQH=Pro+Vt+TT=Adj+N[C,P]+V+TT=Adjsh+N[C,P]"),
        ],
    },
    {
        "num": "Câu 3",
        "vn": "Trẻ em có năng khiếu biết cách điều hướng cảm xúc của mình để hỗ trợ việc học.",
        "en": "Gifted children know how to channel their feelings to assist their learning.",
        "cols": [
            ("Gifted [children]", "Adj+N[C,P]=S"),
            ("[know]", "Vt"),
            ("[how to channel]", "how to V=TT"),
            ("their [feelings]", "Adjsh+N[C,P]=TT"),
            ("[to assist their learning].", "Trạng ngữ=to V+TT=Adjsh+N[U]"),
        ],
    },
    {
        "num": "Câu 4",
        "vn": "Trẻ có năng khiếu vượt trội được hưởng lợi từ sự hỗ trợ phù hợp của người thân.",
        "en": "The very gifted child benefits from appropriate support from close relatives.",
        "cols": [
            ("The very gifted [child]", "Article+Adv=>Adj+N[C,S]=S"),
            ("[benefits]", "Vi"),
            ("[from appropriate support]", "Trạng ngữ=Prep+Adj+N[U]"),
            ("[from close relatives].", "Tính ngữ=Prep+Adj+N[C,P]"),
        ],
    },
    {
        "num": "Câu 5",
        "vn": "Học sinh thực sự xuất sắc đã học hỏi được một lượng kiến thức đáng kể về lĩnh vực của mình.",
        "en": "Really successful students have learnt a considerable amount about their subject.",
        "cols": [
            ("Really successful [students]", "Adv=>Adj+N[C,P]=S"),
            ("[have]", "Aux"),
            ("[learnt]", "Vt"),
            ("a considerable [amount]", "Article+Adj+N[C,S]=TT"),
            ("[about their subject].", "Tính ngữ=Prep+Adjsh+N[C,S]"),
        ],
    },
]


def set_run(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


def add_code_runs(par, code, size):
    """Tách mã hàng 2 thành các run: nhãn chức năng in đậm, loại từ để thường."""
    PLACEHOLDER = ""
    protected = code.replace("=>", PLACEHOLDER)
    runs, token = [], ""
    for ch in protected:
        if ch in "+=":
            if token:
                runs.append(("tok", token)); token = ""
            runs.append(("sep", ch))
        else:
            token += ch
    if token:
        runs.append(("tok", token))
    for kind, text in runs:
        text = text.replace(PLACEHOLDER, "=>")
        bold = kind == "tok" and text in BOLD_TOKENS
        r = par.add_run(text)
        set_run(r, size, bold=bold)


def col_width(c1, c2):
    n = max(len(c1), len(c2))
    return max(MIN_COL, n * DXA_PER_CHAR + 160)


def set_cell_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(w))
    tcW.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Twips(PAGE_W)
sec.page_height = Twips(11906)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Twips(MARGIN))

# Tiêu đề
h = doc.add_heading(level=1)
r = h.add_run("UNIT – GIFTED CHILDREN AND LEARNING (C10T2)")
set_run(r, 16, bold=True)

# Khối ký hiệu — 2 đoạn, in nghiêng
for block in (KY_HIEU_1, KY_HIEU_2):
    p = doc.add_paragraph()
    r = p.add_run(block)
    set_run(r, 9, italic=True)

# Từng câu
for s in SENTENCES:
    widths = [col_width(c1, c2) for c1, c2 in s["cols"]]
    total = sum(widths)
    pt = TABLE_PT
    if total > USABLE:
        scale = USABLE / total
        pt = max(MIN_PT, TABLE_PT * scale)
        widths = [int(w * scale) for w in widths]
        total = sum(widths)

    p = doc.add_paragraph()
    r = p.add_run(s["num"])
    set_run(r, 11, bold=True)

    p = doc.add_paragraph()
    r = p.add_run(s["vn"])
    set_run(r, 11, italic=True)

    p = doc.add_paragraph()
    r = p.add_run(s["en"])
    set_run(r, 11, bold=True)

    table = doc.add_table(rows=2, cols=len(s["cols"]))
    set_table_borders(table)
    table.autofit = False
    for i, ((c1, c2), w) in enumerate(zip(s["cols"], widths)):
        cell1 = table.cell(0, i)
        cell2 = table.cell(1, i)
        set_cell_width(cell1, w)
        set_cell_width(cell2, w)
        par1 = cell1.paragraphs[0]
        par1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par1.add_run(c1)
        set_run(r, pt, bold=True)  # hàng 1: câu tiếng Anh thực tế, đậm toàn bộ
        par2 = cell2.paragraphs[0]
        par2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_code_runs(par2, c2, pt)

    doc.add_paragraph()
    print(f"{s['num']}: {len(s['cols'])} cột, tổng {total} DXA (khả dụng {USABLE}), font {pt}pt")

import os
out = os.environ.get("OUT", "Unit-GiftedChildren-20260710.docx")
doc.save(out)
print("Saved:", out)
