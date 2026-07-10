# -*- coding: utf-8 -*-
"""
Engine dung chung de tao file .docx phan tich PASSAGE theo dinh dang PP1
(khoi bang xep chong, font co dinh de doc, khong ep chu nho).

CACH DUNG (cho moi bai passage):
  1. Viet 1 file data python (vd data_xxx.py) khai bao:
       TITLE          = "PASSAGE ANALYSIS - <TEN BAI>"
       OUT_PATH        = "/duong/dan/toi/file.docx"
       SENTENCES       = [ {"num":..., "vn":..., "en":..., "cols":[(c1,c2), ...]}, ... ]
       EXTRA_KY_HIEU_1 = ""   # (tuy chon) bo sung ky hieu nhom "Chuc nang ngu phap", cach nhau bang " ; "
       EXTRA_KY_HIEU_2 = ""   # (tuy chon) bo sung ky hieu nhom "Loai tu"
       EXTRA_BOLD_TOKENS = [] # (tuy chon) cac nhan CHUC NANG moi can in dam trong hang 2

  2. Chay:  python3 pp1_engine.py data_xxx.py
     -> tao file docx, tu kiem tra cau truc (so cot khop, moi token da khai bao),
        in bao cao PASS/FAIL va thong ke so cau / so khoi bang.

QUY TAC NOI DUNG (bam sat skill unit-builder, Buoc 3):
  - Hang 1 = cum tu tieng Anh that (co ngoac vuong theo quy tac); Hang 2 = ma phan tich.
  - S/BC/TT/TGT/DV: chi ngoac phan dau ngu phap chinh (N / Prep+N / to V); bo nghia dung NGOAI ngoac.
  - Trang ngu, Tinh ngu: ngoac TRON ca cum.
  - MDP/MDQH/MDDT (menh de phu / quan he / danh tu): ngoac TRON CA MENH DE trong 1 cot, hang 2 liet
    ke day du chuoi thanh phan ben trong noi bang "+", mo dau bang nhan MDP=/MDQH=/MDDT=.
  - Dau cau luon dung NGOAI ngoac vuong, gan vao cuoi tu/cum truoc do trong cung 1 o.
  - Nhan CHUC NANG NGU PHAP (S,TT,TGT,BC,DV,Trang ngu,Tinh ngu,MDC,MDP,MDQH,MDDT,CNTT,Conj,Conjp,
    Aux,Advpd,DX,cac ma Verb Vi/Vt/Vikr/Vt[bd]...) -> IN DAM trong hang 2.
    Nhan LOAI TU (Pro,PN,N[...],Article,So tu,Adj*,Adv...,to V,V-ing,V-ed...) -> de thuong.
  - S/BC/TT/TGT/DV viet theo thu tu "cau_tao=NHAN" (nhan dung sau); Trang ngu/Tinh ngu/MDP/MDQH/MDDT/
    CNTT/MDC viet theo thu tu "NHAN=cau_tao" (nhan dung truoc).
  - Neu can ky hieu moi ngoai bang chuan, PHAI khai bao vao EXTRA_KY_HIEU_1/2 truoc khi dung.
"""
import sys
import importlib.util
import re

from docx import Document
from docx.shared import Pt, Twips
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Cambria"
TABLE_PT = 11.0
DXA_PER_CHAR = 135
MIN_COL = 900
PAGE_W = 16838
MARGIN = 500
USABLE = PAGE_W - 2 * MARGIN

BASE_BOLD_TOKENS = {
    "S", "TT", "TGT", "BC", "DV", "Trạng ngữ", "Tính ngữ", "MĐC", "MĐP", "MĐQH", "MĐDT",
    "CNTT", "Conj", "Conjp", "Aux", "Advpđ", "ĐX", "Vlet",
    "Vi", "Vt", "Vikr", "Vtpd", "Vt[bđ]",
}

BASE_KY_HIEU_1 = (
    "Chức năng ngữ pháp: S=Chủ ngữ ; TT=Tân trực tiếp ; TGT=Tân gián tiếp ; BC=Bổ chủ ; DV=Đồng vị ngữ ; "
    "Trạng ngữ=Adv / Adv+N[...] / Prep+N[...] ; Tính ngữ=Prep+N[...] sau danh từ ; "
    "MĐC=Mệnh đề chính ; MĐP=Mệnh đề phụ (câu phụ, do Conjp dẫn) ; MĐQH=Mệnh đề quan hệ (relative clause) ; "
    "MĐDT=Mệnh đề danh từ (noun clause, thường sau that/whether làm tân ngữ, chủ ngữ thật, hoặc bổ nghĩa danh từ) ; "
    "CNTT=Chủ ngữ thật (trong câu chủ ngữ giả “it”) ; Conj=Liên từ đẳng lập (and, but, or...) ; "
    "Conjp=Liên từ / cụm dẫn nhập phụ thuộc (because, when, if, although, as, given that...) ; "
    "Aux=Trợ động từ (will, can, may, must, should, do, have...) ; "
    "Advpđ=Trạng từ phủ định (not, never, no longer, hardly...) ; "
    "ĐX=Thành phần đệm-xen (thán từ, hô ngữ, liên từ chuyển ý: Of course, Moreover...) ; "
    "Vt=Ngoại động từ ; Vi=Nội động từ ; Vikr=Động từ nối (linking verb) ; Vt[bđ]=Ngoại động từ ở thể bị động (be + V3/V-ed)"
)

BASE_KY_HIEU_2 = (
    "Loại từ: Pro=Đại từ (nhân xưng/quan hệ/chỉ định) ; Pro (CN giả)=Đại từ chủ ngữ giả (“it” không mang nghĩa) ; "
    "PN=Danh từ riêng ; N[C,S]=Danh từ đếm được số ít ; N[C,P]=Danh từ đếm được số nhiều ; "
    "N[U]=Danh từ không đếm được ; N1N2=Danh từ kép ; Article=Mạo từ (a/an/the) ; "
    "Số từ=Số đếm ; Adjsh=Tính từ/cấu trúc sở hữu ('s, my...) ; Adjcd=Tính từ chỉ định (this/that/such) ; "
    "Adjbd=Tính từ bất định (some, several, a number of...) ; Adj=Tính từ miêu tả ; Adv=Trạng từ ; Prep=Giới từ ; "
    "Adv[N]=Danh từ chỉ thời gian dùng như trạng ngữ khi không có giới từ đứng trước (vd: Today) ; "
    "V-ing=Danh động từ / phân từ hiện tại dùng làm trạng ngữ giảm lược ; "
    "V-ed=Phân từ quá khứ dùng làm trạng ngữ giảm lược (reduced participle clause) ; "
    "to V=Cụm động từ nguyên mẫu có “to”"
)


def set_run(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


def add_code_runs(par, code, size, bold_tokens):
    PLACEHOLDER = "§"
    protected = code.replace("=>", PLACEHOLDER)
    runs, token = [], ""
    for ch in protected:
        if ch in "+=":
            if token:
                runs.append(("tok", token)); token = ""
            runs.append(("sep", ch))
        else:
            token += ch
    if token:
        runs.append(("tok", token))
    for kind, text in runs:
        text = text.replace(PLACEHOLDER, "=>")
        bold = kind == "tok" and text in bold_tokens
        r = par.add_run(text)
        set_run(r, size, bold=bold)


def col_width(c1, c2):
    n = max(len(c1), len(c2))
    return max(MIN_COL, n * DXA_PER_CHAR + 160)


def set_cell_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(w))
    tcW.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def chunk_columns(cols):
    widths = [col_width(c1, c2) for c1, c2 in cols]
    chunks, current, current_w = [], [], 0
    for (c1, c2), w in zip(cols, widths):
        if current and current_w + w > USABLE:
            chunks.append(current)
            current, current_w = [], 0
        current.append((c1, c2, w))
        current_w += w
    if current:
        chunks.append(current)
    return chunks


def add_block_table(doc, chunk, bold_tokens):
    table = doc.add_table(rows=2, cols=len(chunk))
    set_table_borders(table)
    table.autofit = False
    for i, (c1, c2, w) in enumerate(chunk):
        cell1 = table.cell(0, i)
        cell2 = table.cell(1, i)
        set_cell_width(cell1, w)
        set_cell_width(cell2, w)
        par1 = cell1.paragraphs[0]
        par1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par1.add_run(c1)
        set_run(r, TABLE_PT, bold=True)
        par2 = cell2.paragraphs[0]
        par2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_code_runs(par2, c2, TABLE_PT, bold_tokens)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(1)


def generate_pp1_docx(title, sentences, out_path,
                       extra_ky_hieu_1="", extra_ky_hieu_2="", extra_bold_tokens=None):
    bold_tokens = set(BASE_BOLD_TOKENS)
    if extra_bold_tokens:
        bold_tokens |= set(extra_bold_tokens)
    ky1 = BASE_KY_HIEU_1 + (" ; " + extra_ky_hieu_1 if extra_ky_hieu_1 else "")
    ky2 = BASE_KY_HIEU_2 + (" ; " + extra_ky_hieu_2 if extra_ky_hieu_2 else "")
    sym_text = ky1 + " " + ky2

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Twips(PAGE_W)
    sec.page_height = Twips(11906)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Twips(MARGIN))

    h = doc.add_heading(level=1)
    r = h.add_run(title)
    set_run(r, 15, bold=True)

    for block in (ky1, ky2):
        p = doc.add_paragraph()
        r = p.add_run(block)
        set_run(r, 9, italic=True)

    stats = {"sentences": 0, "blocks": 0, "max_cols": 0}
    for s in sentences:
        p = doc.add_paragraph(); r = p.add_run(s["num"]); set_run(r, 12, bold=True)
        p = doc.add_paragraph(); r = p.add_run(s["vn"]); set_run(r, 11, italic=True)
        p = doc.add_paragraph(); r = p.add_run(s["en"]); set_run(r, 11, bold=True)
        chunks = chunk_columns(s["cols"])
        for chunk in chunks:
            add_block_table(doc, chunk, bold_tokens)
        doc.add_paragraph()
        stats["sentences"] += 1
        stats["blocks"] += len(chunks)
        stats["max_cols"] = max(stats["max_cols"], len(s["cols"]))

    doc.save(out_path)

    # --- tu kiem tra ---
    doc2 = Document(out_path)
    ok = True
    problems = []
    undeclared = set()
    for ti, t in enumerate(doc2.tables, 1):
        r1, r2 = t.rows[0].cells, t.rows[1].cells
        if len(r1) != len(r2):
            ok = False
            problems.append(f"Bang {ti}: lech so cot")
        for c in r2:
            prot = c.text.replace("=>", "§")
            for tok in re.split(r'[+=]', prot):
                tok = tok.replace("§", "=>").strip()
                if tok and "=>" not in tok and tok not in sym_text:
                    undeclared.add(tok)
    if undeclared:
        ok = False
        problems.append(f"Token chua khai bao: {sorted(undeclared)}")

    return ok, problems, stats


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python3 pp1_engine.py <data_file.py>")
        sys.exit(1)
    data_path = sys.argv[1]
    spec = importlib.util.spec_from_file_location("passage_data", data_path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)

    ok, problems, stats = generate_pp1_docx(
        title=mod.TITLE,
        sentences=mod.SENTENCES,
        out_path=mod.OUT_PATH,
        extra_ky_hieu_1=getattr(mod, "EXTRA_KY_HIEU_1", ""),
        extra_ky_hieu_2=getattr(mod, "EXTRA_KY_HIEU_2", ""),
        extra_bold_tokens=getattr(mod, "EXTRA_BOLD_TOKENS", None),
    )
    print(f"Saved: {mod.OUT_PATH}")
    print(f"So cau: {stats['sentences']} | So khoi bang: {stats['blocks']} | Max cot/cau: {stats['max_cols']}")
    if ok:
        print("KIEM TRA: PASS")
    else:
        print("KIEM TRA: FAIL")
        for p in problems:
            print(" -", p)
        sys.exit(2)
